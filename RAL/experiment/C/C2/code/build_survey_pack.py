#!/usr/bin/env python3
"""从无锚定同场录像重新裁切，打成 C2 survey 包 + 20 人打分 Excel。

结构：
  survey/exp1/*.mp4
  survey/exp2/01..10/{token,mishra,xpress,random}XX.mp4 + sentence.txt
  survey/README.md / manifest.json
  survey/score_sheet_20.xlsx

裁切：1920×1080 → 900×1080（x=600..1500），去音轨。
T2 四条件全进卷：token=ours，mishra/xpress=两路话轮级静脸，random=乱序。
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CODE_DIR = Path(__file__).resolve().parent
C2_DIR = CODE_DIR.parent

CROP_W, CROP_H, CROP_X, CROP_Y = 900, 1080, 600, 0

EXP1_MAP = [
    ("T1-01", "happy", "happy01"),
    ("T1-02", "happy", "happy02"),
    ("T1-03", "sad", "sad01"),
    ("T1-04", "sad", "sad02"),
    ("T1-05", "angry", "angry01"),
    ("T1-06", "angry", "angry02"),
    ("T1-07", "fear", "fear01"),
    ("T1-08", "surprise", "surprise01"),
    ("T1-09", "disgust", "disgust01"),
    ("T1-10", "neutral", "neutral01"),
]

EMOTION_ZH = {
    "happy": "快乐",
    "sad": "悲伤",
    "angry": "愤怒",
    "fear": "恐惧",
    "surprise": "惊讶",
    "disgust": "厌恶",
    "neutral": "中性",
}

# T2 对比：token 分别对三条基线
T2_PAIRS = ("mishra", "xpress", "random")


def ffmpeg_crop(src: Path, dst: Path) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "ffmpeg", "-y", "-i", str(src),
        "-vf", f"crop={CROP_W}:{CROP_H}:{CROP_X}:{CROP_Y}",
        "-an", "-c:v", "libx264", "-crf", "18", "-preset", "fast",
        "-pix_fmt", "yuv420p", "-movflags", "+faststart",
        str(dst),
    ]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(f"ffmpeg failed on {src}:\n{r.stderr[-800:]}")


def load_t2_prompts(path: Path) -> dict[str, dict]:
    rows = {}
    with path.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows[row["id"]] = row
    return rows


def write_sentence_txt(path: Path, tid: str, row: dict, folder: str) -> None:
    shown = (row.get("user_prompt") or "").strip()
    text = (
        "【展示给被试的句子】\n"
        f"{shown}\n\n"
        "【仅实验人员查看，勿展示给被试】\n"
        f"user_prompt: {shown}\n"
        f"source_id: {tid}\n"
        f"pattern: {row.get('target_pattern', '')}\n"
        f"files: token{folder}.mp4 / mishra{folder}.mp4 / "
        f"xpress{folder}.mp4 / random{folder}.mp4\n"
        "scoring: 每次二选一选 token 为命中\n"
        "  token_vs_mishra / token_vs_xpress / token_vs_random\n"
        "mishra = Mishra×GPT 话轮级静脸；xpress = Xpress×GPT 话轮级静脸\n"
    )
    path.write_text(text, encoding="utf-8")


def write_readme(path: Path) -> None:
    path.write_text(
        """# C2 人因问卷材料包（给做问卷的同事）

本包只含**已裁切、无声**刺激（20260731 无锚定同场重录）。被试约 20 人；建议先实验 1 再实验 2。

```
survey/
  README.md
  manifest.json
  score_sheet_20.xlsx
  exp1/                      # 情绪识别 10 视频
  exp2/01…10/                # token / mishra / xpress / random + sentence.txt
```

## 共同注意
1. 全部无声；播放时不要出现字幕/文件名。
2. 文件名含条件信息（token01、mishra01…），勿露给被试；打乱题序。
3. 画面约 **900×1080**（x=600..1500，本场重新裁切）。
4. `mishraXX.mp4` / `xpressXX.mp4` 都是**话轮级静脸**（两套 SOTA，标签空间与强度不同）；`randomXX.mp4` 为打乱时间序。

## 实验 1：认表情
七选一：中性 / 快乐 / 惊讶 / 恐惧 / 悲伤 / 厌恶 / 愤怒。
标准答案见 `score_sheet_20.xlsx`「实验1_认表情」。

## 实验 2：哪种脸更贴这句话的情感变化
每句 **三次** 二选一：
- token vs mishra
- token vs xpress
- token vs random
（共 10×3 = 30 题）

先读 `sentence.txt` 里「展示给被试的句子」，再看脸。
问：「哪一张脸更匹配这句话的情感变化？」
计分：选 token 为命中。左右位置见打分表「建议展示」（已奇偶交替）。

## 打分表用法
- 「实验1_认表情」：P01–P20 填中文情绪名。
- 「实验2_二选一」：按建议左右播放，P01–P20 填「左」或「右」。
- 「汇总」：自动统计命中率（含 vs mishra / vs xpress / vs random 分列）。
""",
        encoding="utf-8",
    )


def build_excel(out_xlsx: Path, t2_rows: dict[str, dict]) -> None:
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    lock_fill = PatternFill("solid", fgColor="FFF2CC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws0 = wb.active
    ws0.title = "说明"
    ws0["A1"] = "C2 人因问卷 · 20 人打分记录表（四条件 T2：token/mishra/xpress/random）"
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.merge_cells("A1:F1")
    for i, line in enumerate(
        [
            "1. 先实验1（10 题认表情），再实验2（30 题二选一：10 句 × token对mishra/xpress/random）。",
            "2. 文件名含条件，播放时勿让被试看到；题序建议打乱。",
            "3. 实验1：Pxx 填快乐/悲伤/愤怒/恐惧/惊讶/厌恶/中性。",
            "4. 实验2：按「建议展示」左右顺序播放；Pxx 填「左」或「右」。",
            "5. mishra / xpress 均为话轮级静脸（两套 SOTA）；选 token 计为命中。",
        ],
        start=3,
    ):
        ws0[f"A{i}"] = line
    ws0.column_dimensions["A"].width = 100

    # --- 实验1 ---
    ws1 = wb.create_sheet("实验1_认表情")
    headers1 = [
        "题号", "视频文件", "source_id", "正确答案(中)", "正确答案(英)",
        *[f"P{i:02d}" for i in range(1, 21)],
        *[f"P{i:02d}_命中" for i in range(1, 21)],
    ]
    for c, h in enumerate(headers1, 1):
        cell = ws1.cell(1, c, h)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = thin

    for r, (tid, label, name) in enumerate(EXP1_MAP, start=2):
        for c, v in enumerate(
            [r - 1, f"exp1/{name}.mp4", tid, EMOTION_ZH[label], label], 1
        ):
            cell = ws1.cell(r, c, v)
            cell.alignment = center
            cell.border = thin
            cell.fill = lock_fill
        for p in range(20):
            ac, hc = 6 + p, 26 + p
            ws1.cell(r, ac, "").border = thin
            ws1.cell(r, ac).alignment = center
            ws1.cell(
                r, hc,
                f'=IF({get_column_letter(ac)}{r}="","",'
                f'IF(OR({get_column_letter(ac)}{r}=$D{r},{get_column_letter(ac)}{r}=$E{r}),1,0))',
            ).border = thin
            ws1.cell(r, hc).alignment = center

    for c in range(1, 46):
        ws1.column_dimensions[get_column_letter(c)].width = 12 if c >= 6 else 14
    ws1.column_dimensions["B"].width = 22
    ws1.freeze_panes = "F2"

    # --- 实验2：30 题 ---
    ws2 = wb.create_sheet("实验2_二选一")
    headers2 = [
        "题号", "句子文件夹", "source_id", "pattern", "对比类型",
        "建议左", "建议右", "正确侧(相对建议展示)", "展示句子(摘要)",
        *[f"P{i:02d}" for i in range(1, 21)],
        *[f"P{i:02d}_命中" for i in range(1, 21)],
    ]
    for c, h in enumerate(headers2, 1):
        cell = ws2.cell(1, c, h)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = thin

    row_i, q = 2, 1
    # 记录三类对比所在行，供汇总
    pair_rows = {p: [] for p in T2_PAIRS}
    for idx in range(1, 11):
        folder = f"{idx:02d}"
        tid = f"T2-{idx:02d}"
        row = t2_rows[tid]
        prompt = (row.get("user_prompt") or "").strip()
        summary = prompt if len(prompt) <= 40 else prompt[:37] + "…"
        for other in T2_PAIRS:
            other_name = f"{other}{folder}.mp4"
            token_left = (q % 2 == 1)
            if token_left:
                left, right, correct = f"token{folder}.mp4", other_name, "左"
            else:
                left, right, correct = other_name, f"token{folder}.mp4", "右"
            vals = [
                q, folder, tid, row.get("target_pattern", ""),
                f"token_vs_{other}", left, right, correct, summary,
            ]
            for c, v in enumerate(vals, 1):
                cell = ws2.cell(row_i, c, v)
                cell.alignment = center
                cell.border = thin
                if c <= 9:
                    cell.fill = lock_fill
            for p in range(20):
                ac, hc = 10 + p, 30 + p
                ws2.cell(row_i, ac, "").border = thin
                ws2.cell(row_i, ac).alignment = center
                ws2.cell(
                    row_i, hc,
                    f'=IF({get_column_letter(ac)}{row_i}="","",'
                    f'IF({get_column_letter(ac)}{row_i}=$H{row_i},1,0))',
                ).border = thin
                ws2.cell(row_i, hc).alignment = center
            pair_rows[other].append(row_i)
            row_i += 1
            q += 1

    for c in range(1, 51):
        ws2.column_dimensions[get_column_letter(c)].width = 11 if c >= 10 else 14
    ws2.column_dimensions["I"].width = 36
    ws2.freeze_panes = "J2"

    # --- 汇总 ---
    ws3 = wb.create_sheet("汇总")
    headers3 = [
        "被试", "实验1命中(/10)", "实验1命中率",
        "实验2命中(/30)", "实验2命中率",
        "vs_mishra(/10)", "vs_xpress(/10)", "vs_random(/10)",
    ]
    for c, h in enumerate(headers3, 1):
        cell = ws3.cell(1, c, h)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.border = thin
        cell.alignment = center

    for p in range(20):
        r = p + 2
        pid = f"P{p+1:02d}"
        ws3.cell(r, 1, pid).border = thin
        h1 = get_column_letter(26 + p)
        a1 = get_column_letter(6 + p)
        ws3.cell(r, 2, f"=SUM('实验1_认表情'!{h1}2:{h1}11)").border = thin
        ws3.cell(
            r, 3,
            f'=IF(COUNTA(\'实验1_认表情\'!{a1}2:{a1}11)=0,"",B{r}/10)',
        ).border = thin
        h2 = get_column_letter(30 + p)
        a2 = get_column_letter(10 + p)
        ws3.cell(r, 4, f"=SUM('实验2_二选一'!{h2}2:{h2}31)").border = thin
        ws3.cell(
            r, 5,
            f'=IF(COUNTA(\'实验2_二选一\'!{a2}2:{a2}31)=0,"",D{r}/30)',
        ).border = thin
        for col, other in enumerate(T2_PAIRS, start=6):
            cells = "+".join(f"'实验2_二选一'!{h2}{rr}" for rr in pair_rows[other])
            ws3.cell(r, col, f"={cells}").border = thin
        for c in range(1, 9):
            ws3.cell(r, c).alignment = center

    r = 22
    ws3.cell(r, 1, "平均").font = Font(bold=True)
    for c, col in enumerate("BCDEFGH", start=2):
        ws3.cell(r, c, f"=AVERAGE({col}2:{col}21)").border = thin
        ws3.cell(r, c).alignment = center
    ws3.cell(r, 1).border = thin

    for c, w in enumerate([8, 14, 12, 14, 12, 14, 14, 14], 1):
        ws3.column_dimensions[get_column_letter(c)].width = w

    wb.save(out_xlsx)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--session", type=Path,
        default=C2_DIR / "result" / "session_noanchor_20260731",
    )
    ap.add_argument("--out-dir", type=Path, default=C2_DIR / "output" / "survey")
    ap.add_argument(
        "--zip-path", type=Path,
        default=C2_DIR / "output" / "C2_survey_noanchor_20260731.zip",
    )
    args = ap.parse_args()

    session: Path = args.session
    out: Path = args.out_dir
    if out.exists():
        shutil.rmtree(out)
    (out / "exp1").mkdir(parents=True)

    t2_rows = load_t2_prompts(C2_DIR / "data" / "t2_prompts.csv")

    exp1_manifest = []
    for tid, label, name in EXP1_MAP:
        src = session / "t1" / "ours" / tid / "video.mp4"
        dst = out / "exp1" / f"{name}.mp4"
        print(f"[exp1] {tid} -> {dst.name}", flush=True)
        ffmpeg_crop(src, dst)
        exp1_manifest.append(
            {"survey_file": f"exp1/{name}.mp4", "name": name, "label": label, "source_id": tid}
        )

    exp2_manifest = []
    cond_src = {
        "token": "ours",
        "mishra": "mishra_gpt",
        "xpress": "xpress_gpt",
        "random": "random",
    }
    for idx in range(1, 11):
        folder = f"{idx:02d}"
        tid = f"T2-{idx:02d}"
        d = out / "exp2" / folder
        d.mkdir(parents=True)
        for survey_name, src_cond in cond_src.items():
            src = session / "t2" / src_cond / tid / "video.mp4"
            name = f"{survey_name}{folder}.mp4"
            print(f"[exp2/{folder}] {src_cond} -> {name}", flush=True)
            ffmpeg_crop(src, d / name)
        write_sentence_txt(d / "sentence.txt", tid, t2_rows[tid], folder)
        exp2_manifest.append(
            {
                "folder": folder,
                "source_id": tid,
                "pattern": t2_rows[tid].get("target_pattern", ""),
                "conditions": list(cond_src.keys()),
            }
        )

    write_readme(out / "README.md")
    (out / "manifest.json").write_text(
        json.dumps(
            {
                "session": session.name,
                "crop": f"{CROP_W}x{CROP_H} x={CROP_X}..{CROP_X+CROP_W}",
                "silent": True,
                "exp2_conditions": list(cond_src.keys()),
                "exp2_pairs": [f"token_vs_{p}" for p in T2_PAIRS],
                "exp1": exp1_manifest,
                "exp2": exp2_manifest,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    build_excel(out / "score_sheet_20.xlsx", t2_rows)

    args.zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(args.zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(out.rglob("*")):
            if f.is_file():
                zf.write(f, arcname=str(Path("survey") / f.relative_to(out)))
    print(f"Wrote {out}")
    print(f"Wrote {args.zip_path} ({args.zip_path.stat().st_size/1e6:.1f} MB)")


if __name__ == "__main__":
    main()
